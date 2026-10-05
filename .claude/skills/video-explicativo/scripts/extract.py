#!/usr/bin/env python3
"""Extrae texto plano de PDF, DOCX, PPTX, TXT, MD, HTML o CSV.
Uso: python3 extract.py archivo [--out texto.txt]"""
import re, subprocess, sys
from pathlib import Path


def need(pkg, mod=None):
    try:
        __import__(mod or pkg)
    except ImportError:
        subprocess.run([sys.executable, "-m", "pip", "install", "-q", pkg], check=True)


def extract(p: Path) -> str:
    ext = p.suffix.lower()
    if ext == ".pdf":
        need("pypdf")
        from pypdf import PdfReader
        return "\n\n".join((pg.extract_text() or "") for pg in PdfReader(str(p)).pages)
    if ext == ".docx":
        need("python-docx", "docx")
        import docx
        d = docx.Document(str(p))
        parts = [para.text for para in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext == ".pptx":
        need("python-pptx", "pptx")
        from pptx import Presentation
        out = []
        for i, s in enumerate(Presentation(str(p)).slides, 1):
            out.append(f"--- Diapositiva {i} ---")
            for sh in s.shapes:
                if sh.has_text_frame:
                    out.append(sh.text_frame.text)
        return "\n".join(out)
    if ext in (".html", ".htm"):
        t = p.read_text(encoding="utf-8", errors="ignore")
        t = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", t, flags=re.S | re.I)
        return re.sub(r"<[^>]+>", " ", t)
    return p.read_text(encoding="utf-8", errors="ignore")


if __name__ == "__main__":
    src = Path(sys.argv[1])
    txt = re.sub(r"\n{3,}", "\n\n", extract(src)).strip()
    if "--out" in sys.argv:
        Path(sys.argv[sys.argv.index("--out") + 1]).write_text(txt, encoding="utf-8")
    words = len(txt.split())
    print(txt if "--out" not in sys.argv else f"{words} palabras extraídas")
